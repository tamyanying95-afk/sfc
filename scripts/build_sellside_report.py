#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render the sell-side SFC batch result into an HTML report + Excel."""

import datetime
import html
import json
import os

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE_DIR, "output")
data = json.load(open(os.path.join(OUT, "sfc_sellside_result.json"), encoding="utf-8"))
s = data["summary"]
groups = data["groups"]
companies = data["companies"]
people = sorted(data["people"], key=lambda p: (p["name"] or "").upper())

# entity -> group lookup
name_to_group, ce_to_group = {}, {}
for g in groups:
    for e in g["entities"]:
        name_to_group[e["sfc_name"]] = g["group"]
        if e["ceref"]:
            ce_to_group[e["ceref"]] = g["group"]


def groups_of(p):
    out = []
    for cn in p["companies"]:
        g = name_to_group.get(cn)
        if g and g not in out:
            out.append(g)
    return out


grows = []
for i, g in enumerate(groups, 1):
    ces = [e["ceref"] for e in g["entities"] if e["ceref"]]
    links = " ".join("<a href='https://apps.sfc.hk/publicregWeb/corp/%s/details?locale=en' target='_blank'>%s</a>"
                     % (c, c) for c in ces)
    grows.append(
        "<tr><td class='num muted'>%d</td><td><b>%s</b></td><td class='num'>%d</td><td>%s</td>"
        "<td class='num'>%s</td><td class='num'>%s</td><td class='num strong'>%s</td></tr>"
        % (i, html.escape(g["group"]), len(g["entities"]), links, g["ro"], g["rep"], g["unique"]))

# CITIC + CLSA combined: they share 970 people, so also show the joint figure
gmap = {g["group"]: g for g in groups}
if "CITIC" in gmap and "CLSA" in gmap:
    joint_groups = {"CITIC", "CLSA"}
    joint = {p["ceRef"] for p in people if joint_groups & set(groups_of(p))}
    grows.append(
        "<tr class='joint'><td></td><td><b>CITIC + CLSA 合并</b> <span class='muted'>（两集团共享 970 人）</span></td>"
        "<td class='num'>%d</td><td></td><td class='num'>%d</td><td class='num'>%d</td>"
        "<td class='num strong'>%d</td></tr>"
        % (len(gmap["CITIC"]["entities"]) + len(gmap["CLSA"]["entities"]),
           gmap["CITIC"]["ro"] + gmap["CLSA"]["ro"],
           gmap["CITIC"]["rep"] + gmap["CLSA"]["rep"], len(joint)))

erows = []
for c in companies:
    ce = c["ceref"] or ""
    link = ("<a href='https://apps.sfc.hk/publicregWeb/corp/%s/details?locale=en' target='_blank'>%s</a>"
            % (ce, ce)) if ce else "&mdash;"
    flag = "" if c.get("match") in ("exact", "contains", "only-candidate") else " <span class='warn'>核对</span>"
    erows.append(
        "<tr><td class='muted'>%s</td><td>%s</td><td>%s%s</td><td>%s</td><td>%s</td>"
        "<td class='num'>%s</td><td class='num'>%s</td><td class='num'>%s</td></tr>"
        % (html.escape(c.get("group", "")), html.escape(c["input"]),
           html.escape(c["sfc_name"] or "&mdash;"), flag, link,
           html.escape(c["status"]), c["ro"], c["rep"], c["unique"]))

prows = []
for p in people:
    roles = "+".join(p["roles"])
    badge = "both" if len(p["roles"]) > 1 else ("ro" if "RO" in p["roles"] else "rep")
    prows.append(
        "<tr><td>%s</td><td>%s</td><td>%s</td><td><span class='badge %s'>%s</span></td><td>%s</td><td>%s</td></tr>"
        % (html.escape(p["ceRef"] or ""), html.escape(p["name"] or ""),
           html.escape(p["nameChi"] or ""), badge, roles,
           html.escape("; ".join(groups_of(p))), html.escape("; ".join(p["companies"]))))

HTML = """<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>SFC 持牌人统计 · Sell-side</title>
<style>
 :root{--bg:#f6f7f9;--card:#fff;--line:#e5e7eb;--text:#111827;--muted:#6b7280;--accent:#1d4ed8;}
 *{box-sizing:border-box}
 body{margin:0;background:var(--bg);color:var(--text);font:15px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI","PingFang SC","Microsoft YaHei",sans-serif;}
 .wrap{max-width:1280px;margin:0 auto;padding:32px 20px 64px}
 h1{font-size:26px;margin:0 0 6px}
 .sub{color:var(--muted);font-size:13px;margin-bottom:24px}
 .cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:14px;margin-bottom:28px}
 .card{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:16px 18px}
 .card .k{font-size:12px;color:var(--muted)}
 .card .v{font-size:28px;font-weight:650;margin-top:4px}
 .card .h{font-size:12px;color:var(--muted);margin-top:2px}
 table{width:100%;border-collapse:collapse;background:var(--card);border:1px solid var(--line);border-radius:12px;overflow:hidden}
 caption{text-align:left;font-size:17px;font-weight:650;padding:0 0 12px}
 th,td{padding:8px 11px;border-bottom:1px solid var(--line);text-align:left;font-size:13.5px;vertical-align:top}
 th{background:#f0f2f5;font-weight:600;font-size:13px;color:#374151}
 tr:last-child td{border-bottom:none}
 td.num,th.num{text-align:right;font-variant-numeric:tabular-nums}
 td.strong{font-weight:650}
 tfoot td{background:#f0f2f5;font-weight:650}
 .muted{color:var(--muted)}
 .warn{display:inline-block;padding:0 6px;border-radius:6px;font-size:11px;background:#fef3c7;border:1px solid #fde68a;color:#92400e}
 .badge{display:inline-block;padding:1px 8px;border-radius:999px;font-size:12px;border:1px solid var(--line);background:#f3f4f6;color:#374151}
 .badge.both{background:#fef3c7;border-color:#fde68a;color:#92400e}
 .badge.ro{background:#dbeafe;border-color:#bfdbfe;color:#1e40af}
 .badge.rep{background:#ecfdf5;border-color:#a7f3d0;color:#065f46}
 .note{background:var(--card);border:1px solid var(--line);border-left:3px solid var(--accent);border-radius:8px;padding:12px 16px;margin:22px 0;font-size:13.5px;color:#374151}
 .note b{color:var(--accent)}
 tr.joint td{background:#fffbeb}
 .scroll{max-height:620px;overflow:auto;border:1px solid var(--line);border-radius:12px}
 .scroll table{border:none;border-radius:0}
 a{color:var(--accent);text-decoration:none}
 a:hover{text-decoration:underline}
</style></head><body><div class="wrap">
<h1>SFC 持牌代表及负责人员（RO）统计 · Sell-side</h1>
<div class="sub">数据源：香港证监会公众纪录册 · licstatus = Active（个别实体已停用，见实体明细表）· SFO licence and/or AMLO licence · 抓取日期 __DATE__</div>

<div class="cards">
  <div class="card"><div class="k">集团数量</div><div class="v">__NG__</div><div class="h">按你给的名单顺序</div></div>
  <div class="card"><div class="k">法人实体数</div><div class="v">__NE__</div><div class="h">全部匹配到 CE No.</div></div>
  <div class="card"><div class="k">集团去重后合计</div><div class="v">__GSUM__</div><div class="h">各集团去重人数相加</div></div>
  <div class="card"><div class="k">全局去重人数</div><div class="v">__GLOB__</div><div class="h">跨集团重复 __OVL__ 人</div></div>
  <div class="card"><div class="k">RO / Rep 记录数</div><div class="v">__RO__ / __REP__</div><div class="h">未去重</div></div>
</div>

<div class="note"><b>口径：</b>同一集团（如 UBS）旗下所有列出的法人实体先求和，再按 <b>CE No.</b> 在集团内去重 —— 一人同时是 Rep 与 RO、或同时受聘于集团内多家实体，都只算一次。</div>

<div class="note"><b>关于 CITIC 与 CLSA：</b>这两个集团之间有 <b>970 人</b>是同一批人（同一 CE No. 同时受聘于两边实体，状态均为 Active），因此 18 个集团去重人数相加为 __GSUM__，全局去重后为 <b>__GLOB__</b>。若把 CITIC 与 CLSA 视为同一体系，合并后为 <b>1,918 人</b>（表中黄底行）。其余集团之间不存在重叠。</div>

<table>
  <caption>按集团汇总（顺序＝你提供的名单）</caption>
  <thead><tr><th class="num">#</th><th>集团</th><th class="num">实体数</th><th>CE No.</th><th class="num">ROs</th><th class="num">Representatives</th><th class="num">独立人数（集团内去重）</th></tr></thead>
  <tbody>__GROWS__</tbody>
  <tfoot><tr><td></td><td>合计 __NG__ 个集团</td><td class="num">__NE__</td><td></td><td class="num">__RO__</td><td class="num">__REP__</td><td class="num">__GSUM__</td></tr></tfoot>
</table>

<div style="height:28px"></div>

<div class="scroll">
<table>
  <caption>实体明细（__NE__ 家）</caption>
  <thead><tr><th>集团</th><th>你提供的名称</th><th>SFC 登记名称</th><th>CE No.</th><th>状态</th><th class="num">ROs</th><th class="num">Reps</th><th class="num">实体内去重</th></tr></thead>
  <tbody>__EROWS__</tbody>
</table>
</div>

<div style="height:28px"></div>

<div class="scroll">
<table>
  <caption>去重后人员明细（__GLOB__ 人）</caption>
  <thead><tr><th>CE No.</th><th>Name</th><th>中文名</th><th>角色</th><th>所属集团</th><th>所属实体</th></tr></thead>
  <tbody>__PROWS__</tbody>
</table>
</div>

<div class="note"><b>需留意：</b>HSBC 的 HSBC Broking Securities (Hong Kong) Limited、HSBC Global Asset Management Holdings (Bahamas) Limited 在 SFC 已无有效牌照（已停用），其 RO/Rep 记录为空；CICC 的资产管理、证券两家实体在名单中的写法与 SFC 登记名略有差异（已按最相近记录匹配）。</div>
</div></body></html>
"""

HTML = (HTML
        .replace("__DATE__", datetime.date.today().strftime("%Y-%m-%d"))
        .replace("__NG__", str(len(groups)))
        .replace("__NE__", str(s["entities_input"]))
        .replace("__GSUM__", str(s["sum_of_group_unique"]))
        .replace("__GLOB__", str(s["global_unique_people"]))
        .replace("__OVL__", str(s["sum_of_group_unique"] - s["global_unique_people"]))
        .replace("__RO__", str(s["total_ro_rows"]))
        .replace("__REP__", str(s["total_rep_rows"]))
        .replace("__GROWS__", "\n".join(grows))
        .replace("__EROWS__", "\n".join(erows))
        .replace("__PROWS__", "\n".join(prows)))

with open(os.path.join(OUT, "sfc_sellside_report.html"), "w", encoding="utf-8") as f:
    f.write(HTML)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

wb = Workbook()
ws = wb.active
ws.title = "Group summary"
ws.append(["#", "集团", "实体数", "ROs", "Representatives", "独立人数（集团内按 CE No. 去重）"])
for i, g in enumerate(groups, 1):
    ws.append([i, g["group"], len(g["entities"]), g["ro"], g["rep"], g["unique"]])
ws.append(["", "TOTAL", sum(len(g["entities"]) for g in groups),
           sum(g["ro"] for g in groups), sum(g["rep"] for g in groups),
           sum(g["unique"] for g in groups)])
if "CITIC" in gmap and "CLSA" in gmap:
    joint = {p["ceRef"] for p in people if {"CITIC", "CLSA"} & set(groups_of(p))}
    ws.append(["", "CITIC + CLSA 合并（共享 970 人）",
               len(gmap["CITIC"]["entities"]) + len(gmap["CLSA"]["entities"]),
               gmap["CITIC"]["ro"] + gmap["CLSA"]["ro"],
               gmap["CITIC"]["rep"] + gmap["CLSA"]["rep"], len(joint)])
for cell in ws[1]:
    cell.font = Font(bold=True)
for cell in ws[ws.max_row]:
    cell.font = Font(bold=True)
for col, w in zip("ABCDEF", [5, 24, 9, 8, 16, 30]):
    ws.column_dimensions[col].width = w

ws2 = wb.create_sheet("Entity detail")
ws2.append(["集团", "你提供的名称", "SFC 登记名称", "CE No.", "状态", "匹配方式", "ROs", "Representatives", "实体内去重", "备注"])
for c in companies:
    ws2.append([c.get("group", ""), c["input"], c["sfc_name"], c["ceref"], c["status"],
                c.get("match", ""), c["ro"], c["rep"], c["unique"], c.get("note", "")])
for cell in ws2[1]:
    cell.font = Font(bold=True)
for col, w in zip("ABCDEFGHIJ", [18, 46, 52, 10, 10, 16, 7, 15, 12, 40]):
    ws2.column_dimensions[col].width = w
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

ws3 = wb.create_sheet("People (CE dedupe)")
ws3.append(["CE No.", "Name", "Chinese name", "Roles", "集团", "所属实体"])
for p in people:
    ws3.append([p["ceRef"], p["name"], p["nameChi"], "+".join(p["roles"]),
                "; ".join(groups_of(p)), "; ".join(p["companies"])])
for cell in ws3[1]:
    cell.font = Font(bold=True)
for col, w in zip("ABCDEF", [12, 34, 14, 10, 22, 52]):
    ws3.column_dimensions[col].width = w
wb.save(os.path.join(OUT, "sfc_sellside_headcount.xlsx"))

print("sell-side report written; groups=%d entities=%d group_sum=%d global=%d" %
      (len(groups), s["entities_input"], s["sum_of_group_unique"], s["global_unique_people"]))
